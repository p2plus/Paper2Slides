"""Exercise uploads, regeneration, status and results with real input files.

The pipeline boundary is replaced to avoid network/model calls. Input selection,
parsing, session handling, state persistence and API responses remain real.
"""

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from api import server
from paper2slides.core.state import load_state, save_state
from paper2slides.document_parser import parse_native_document


def docx_bytes():
    document = Document()
    document.add_heading("Word upload", 1)
    document.add_paragraph("Body reaches the pipeline.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture
def api_client(tmp_path, monkeypatch):
    uploads = tmp_path / "uploads"
    outputs = tmp_path / "outputs"
    uploads.mkdir()
    outputs.mkdir()
    monkeypatch.setattr(server, "UPLOAD_DIR", uploads)
    monkeypatch.setattr(server, "OUTPUT_DIR", outputs)
    monkeypatch.setattr(server, "session_manager", server.SessionManager())
    monkeypatch.setattr(server.app.state, "results", {}, raising=False)
    calls = []

    async def pipeline(base_dir, config_dir, config, from_stage, session_id, manager):
        calls.append(config)
        state = load_state(config_dir)
        state["stages"]["rag"] = "running"
        save_state(config_dir, state)
        assert config["skip_parsing"] is False
        source = Path(config["input_path"])
        paths = [source] if source.is_file() else sorted(source.iterdir())
        for path in paths:
            if path.suffix.lower() in {".md", ".markdown", ".txt", ".docx"}:
                parse_native_document(path, base_dir / "rag_output")
        # Stand in for successful summary, planning and image generation.
        result_dir = config_dir / "20261002_120000"
        result_dir.mkdir(exist_ok=True)
        (result_dir / "slide.png").write_bytes(b"generated image")
        (result_dir / "slides.pdf").write_bytes(b"generated slides")
        state["stages"] = {stage: "completed" for stage in state["stages"]}
        save_state(config_dir, state)

    monkeypatch.setattr(server, "run_pipeline", pipeline)
    with TestClient(server.app) as client:
        yield client, calls


@pytest.mark.parametrize(
    "filename,body",
    [
        ("notes.md", b"# Markdown upload\n\nBody reaches the pipeline."),
        ("notes.MARKDOWN", b"# Long Markdown extension"),
        ("report.DOCX", None),
        ("paper.PDF", b"%PDF-1.4"),
    ],
)
@pytest.mark.parametrize("output_type", ["slides", "poster"])
def test_upload_status_result_and_regeneration(api_client, filename, body, output_type):
    client, calls = api_client
    response = client.post(
        "/api/chat",
        data={"style": "academic", "output_type": output_type},
        files={"files": (filename, body if body is not None else docx_bytes())},
    )
    assert response.status_code == 200
    session_id = response.json()["session_id"]
    assert calls[0]["input_path"].endswith(filename)
    status = client.get(f"/api/status/{session_id}").json()
    assert status["status"] == "completed"
    result = client.get(f"/api/result/{session_id}")
    assert result.status_code == 200
    assert len(result.json()["slides"]) == 1
    assert result.json()[
        "ppt_url" if output_type == "slides" else "poster_url"
    ].endswith("slides.pdf")
    regenerated = client.post(
        "/api/chat",
        data={
            "session_id": session_id,
            "style": "academic",
            "output_type": output_type,
        },
    )
    assert regenerated.status_code == 200
    assert len(calls) == 2
    assert calls[1]["input_path"] == calls[0]["input_path"]


def test_mixed_uploads_keep_all_supported_inputs(api_client):
    client, calls = api_client
    response = client.post(
        "/api/chat",
        data={"style": "academic"},
        files=[
            ("files", ("notes.md", b"# Markdown source")),
            ("files", ("notes.docx", docx_bytes())),
            ("files", ("paper.pdf", b"%PDF-1.4")),
        ],
    )
    session_id = response.json()["session_id"]
    assert len(calls[0]["pdf_paths"]) == 3
    assert Path(calls[0]["input_path"]).name == session_id
    assert client.get(f"/api/status/{session_id}").json()["status"] == "completed"
    assert client.get(f"/api/result/{session_id}").status_code == 200
    parsed_files = list(server.OUTPUT_DIR.rglob("*.md"))
    assert len(parsed_files) == 2


@pytest.mark.parametrize("message", ["", "Use a clean blue layout"])
@pytest.mark.parametrize(
    "filename,body", [("empty.md", b""), ("broken.docx", b"invalid")]
)
def test_parse_failure_is_persisted_and_visible_in_status(
    api_client, filename, body, message
):
    client, _ = api_client
    response = client.post(
        "/api/chat",
        data={"style": "academic", "message": message},
        files={"files": (filename, body)},
    )
    session_id = response.json()["session_id"]
    status = client.get(f"/api/status/{session_id}").json()
    assert status["status"] == "failed"
    assert filename in status["error"]
    assert client.get(f"/api/result/{session_id}").status_code == 500
    assert client.get("/api/session/running").json()["has_running_session"] is False
