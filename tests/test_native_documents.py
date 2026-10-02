"""Issue #30: real document fixtures, without Office or OCR executables."""

import asyncio
import json
import sys
from io import BytesIO
from pathlib import Path
from types import ModuleType

import pytest
from docx import Document
from helpers import load_batch_parser_module, load_parser_module
from PIL import Image

from paper2slides.document_parser import parse_native_document


@pytest.mark.parametrize("extension", ["md", "MD", "markdown", "txt"])
def test_text_keeps_unicode_and_markdown(tmp_path, extension):
    source = tmp_path / f"notes.{extension}"
    text = "# Résumé — 연구\n\n**Results**: x < y & z > 0.\n\n```python\nprint('hello')\n```\n"
    source.write_text(text, encoding="utf-8-sig")
    blocks = parse_native_document(source, tmp_path / "output")
    markdown = next((tmp_path / "output").rglob("*.md"))
    assert markdown.read_text(encoding="utf-8") == text
    assert blocks == [{"type": "text", "text": text, "page_idx": 0}]


def test_markdown_keeps_local_images_when_moved(tmp_path):
    source = tmp_path / "notes.md"
    image = tmp_path / "diagram.png"
    Image.new("RGB", (10, 10), "red").save(image)
    source.write_text("# Notes\n\n![Chart](diagram.png)\n", encoding="utf-8")
    parse_native_document(source, tmp_path / "output")
    markdown = next((tmp_path / "output").rglob("*.md"))
    assert "![Chart](images/" in markdown.read_text()
    assert next(markdown.parent.rglob("*.png")).read_bytes() == image.read_bytes()


def make_docx(source):
    document = Document()
    document.add_heading("Research résumé", 1)
    document.add_paragraph("Results: α < β & γ > 0.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "95%"
    image = BytesIO()
    Image.new("RGB", (10, 10), "blue").save(image, format="PNG")
    image.seek(0)
    document.add_picture(image)
    document.add_paragraph("Conclusion after the figure.")
    document.save(source)


@pytest.mark.parametrize("parser_type", ["mineru", "docling"])
def test_docx_keeps_order_tables_and_images_without_subprocess(
    tmp_path, monkeypatch, parser_type
):
    module = load_parser_module(f"native_{parser_type}_under_test")
    monkeypatch.setattr(
        module.subprocess,
        "run",
        lambda *a, **kw: pytest.fail("No PDF conversion expected"),
    )
    parser = (
        module.MineruParser() if parser_type == "mineru" else module.DoclingParser()
    )
    source = tmp_path / "report.DOCX"
    make_docx(source)
    blocks = parser.parse_document(source, output_dir=tmp_path / "output")
    assert [block["type"] for block in blocks] == [
        "text",
        "text",
        "table",
        "image",
        "text",
    ]
    assert blocks[0]["text"] == "# Research résumé"
    assert "Accuracy" in blocks[2]["table_body"]
    assert Path(blocks[3]["img_path"]).is_file()
    markdown = next((tmp_path / "output").rglob("*.md"))
    content = markdown.read_text(encoding="utf-8")
    assert (
        content.index("Accuracy") < content.index("![](") < content.index("Conclusion")
    )
    assert json.loads(next((tmp_path / "output").rglob("*.json")).read_text()) == blocks


@pytest.mark.parametrize("parser_type", ["mineru", "docling"])
def test_batch_parses_markdown_and_docx_with_identical_stems(tmp_path, parser_type):
    source = tmp_path / "sources"
    source.mkdir()
    (source / "notes.markdown").write_text("# Vault note\n\nMarkdown body.")
    make_docx(source / "notes.docx")
    batch = load_batch_parser_module().BatchParser(
        parser_type=parser_type,
        show_progress=False,
        skip_installation_check=True,
        skip_parsing=False,
    )
    result = batch.process_batch([str(source)], str(tmp_path / "output"))
    assert result.total_files == 2
    assert len(result.successful_files) == 2
    assert result.failed_files == []
    assert len(list((tmp_path / "output").rglob("*.md"))) == 2


@pytest.mark.parametrize(
    "extension,body,error",
    [
        ("md", b" \n", "Document is empty"),
        ("docx", b"not a zip", "Cannot read Word document"),
        ("md", b"\xff\xfeinvalid", "save the document as UTF-8"),
    ],
)
def test_invalid_documents_report_filename(tmp_path, extension, body, error):
    source = tmp_path / f"broken.{extension}"
    source.write_bytes(body)
    with pytest.raises(ValueError, match=error) as raised:
        parse_native_document(source, tmp_path / "output")
    assert source.name in str(raised.value)


def test_fast_mode_stops_if_one_mixed_upload_fails(tmp_path, monkeypatch):
    from paper2slides.core.stages.rag_stage import run_rag_stage

    source = tmp_path / "sources"
    source.mkdir()
    (source / "notes.md").write_text("# Valid input")
    (source / "broken.docx").write_bytes(b"not a Word file")
    monkeypatch.setitem(
        sys.modules, "paper2slides.raganything.batch_parser", load_batch_parser_module()
    )
    rag = ModuleType("paper2slides.rag")
    rag.RAG_PAPER_QUERIES = {}
    monkeypatch.setitem(sys.modules, "paper2slides.rag", rag)
    with pytest.raises(ValueError, match="broken.docx"):
        asyncio.run(
            run_rag_stage(
                tmp_path / "project",
                {
                    "input_path": str(source),
                    "fast_mode": True,
                    "content_type": "paper",
                },
            )
        )
